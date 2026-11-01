"""Regenerate the UK behavioural rates, prices and value-pool entries.

Sources, all byte-stable files:

* ONS time-series JSON endpoints (``/timeseries/<cdid>/<dataset>/previous/v<n>/data``),
  pinned to a published version so the bytes never move: internet sales share,
  average weekly earnings, vacancies, redundancies, unemployment, employment,
  inactivity, claimant count, household saving ratio and GDP growth. The
  profitability series have no pinned version carrying 2025, so they read the
  current endpoint.
* The ONS consumer price inflation time-series CSV (MM23, pinned version) for
  the RPI average prices of selected items.
* The ONS business demography workbook for births, deaths, active enterprises
  and survival rates.
* The DESNZ weekly road fuel prices CSV published on GOV.UK.
* Ofgem's price cap summary documents for the default tariff cap levels.

Run from the repository root:

    uv run python tools/build_gb_behaviour.py
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import openpyxl
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import evidence, fetch, write  # noqa: E402

ONS = "https://www.ons.gov.uk"
ONS_PUBLISHER = "Office for National Statistics"
OGL = "OGL-UK-3.0"
ONS_LICENSE_URL = "https://api.beta.ons.gov.uk/v1/datasets/cpih01"
YEAR = 2025
QUARTERS = 24
MONTHS = {m: f"{i:02d}" for i, m in enumerate(["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], start=1)}
QUARTER_END = {"Q1": "03-31", "Q2": "06-30", "Q3": "09-30", "Q4": "12-31"}
QUARTER_START = {"Q1": "01-01", "Q2": "04-01", "Q3": "07-01", "Q4": "10-01"}

MM23_URL = f"{ONS}/file?uri=/economy/inflationandpriceindices/datasets/consumerpriceindices/current/previous/v138/mm23.csv"
DEMOGRAPHY_URL = f"{ONS}/file?uri=/businessindustryandtrade/business/activitysizeandlocation/datasets/businessdemographyreferencetable/current/businessdemographyexceltables2024.xlsx"
FUEL_CSV_URL = "https://assets.publishing.service.gov.uk/media/6adf6252d798ae11e9bc5d07/CSV__2018_-__.csv"
FUEL_API_URL = "https://www.gov.uk/api/content/government/statistics/weekly-road-fuel-prices"
OFGEM_FILES = "https://www.ofgem.gov.uk/sites/default/files"
OFGEM_COPYRIGHT_URL = "https://www.ofgem.gov.uk/copyright"
OFGEM_PUBLISHER = "Ofgem (Office of Gas and Electricity Markets)"
DESNZ_PUBLISHER = "Department for Energy Security and Net Zero (GOV.UK)"


# ---------------------------------------------------------------- shared helpers


def collapse(text: str) -> str:
    return " ".join(text.split())


def resource_text(raw: bytes) -> str:
    """The searchable text of a fetched resource, mirroring the verifier: PDF pages, ZIP text members, or bytes."""
    if raw.startswith(b"%PDF"):
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    if raw.startswith(b"PK\x03\x04"):
        archive = zipfile.ZipFile(io.BytesIO(raw))
        return "\n".join(archive.read(i).decode("utf-8", errors="replace") for i in archive.infolist() if i.filename.lower().endswith(".xml"))
    return raw.decode("utf-8", errors="replace")


def cite(publisher: str, title: str, url: str, raw: bytes, locator: str, value: Any, unit: str, snippet: str, license_id: str, role: str = "primary") -> dict:
    """Evidence item that is checked against the fetched bytes at build time, so a bad snippet fails here rather than in the verifier."""
    text = collapse(resource_text(raw))
    for needle, what in ((snippet, "snippet"), (str(value), "value")):
        if collapse(needle) not in text:
            raise ValueError(f"{what} {needle!r} not found in {url}")
    if len(snippet.split()) > 24:
        raise ValueError(f"snippet exceeds 24 words: {snippet!r}")
    return evidence(publisher, title, url, locator, value, unit, snippet, license_id, raw, role)


def ons_license_item(raw: bytes) -> dict:
    return cite(ONS_PUBLISHER, "ONS dataset metadata cpih01 (licence statement for ONS datasets)", ONS_LICENSE_URL, raw, "license", "Open Government Licence v3.0", "text", '"license":"Open Government Licence v3.0"', OGL, role="supporting")


# ---------------------------------------------------------------- ONS time series


class OnsSeries:
    """One ONS time series read from its JSON endpoint, optionally pinned to a published version."""

    def __init__(self, path: str, cdid: str, dataset: str, version: int | None = None) -> None:
        self.cdid, self.dataset = cdid.upper(), dataset.upper()
        suffix = f"/previous/v{version}" if version else ""
        self.url = f"{ONS}/{path}/timeseries/{cdid.lower()}/{dataset.lower()}{suffix}/data"
        self.raw = fetch(self.url)
        self.data = json.loads(self.raw)
        self.description = self.data["description"]
        self.version = version

    @property
    def title(self) -> str:
        return self.description["title"]

    def months(self, year: int) -> list[tuple[str, float, dict]]:
        return [(f"{p['year']}-{MONTHS[p['date'].split()[1]]}", float(p["value"]), p) for p in self.data["months"] if p["year"] == str(year)]

    def quarters(self, count: int) -> list[tuple[str, float, dict]]:
        """The last ``count`` quarters that carry a value; ONS leaves the value empty where a series has not been compiled."""
        observed = [p for p in self.data["quarters"] if p["value"] != ""]
        return [(f"{p['year']}-{p['quarter']}", float(p["value"]), p) for p in observed[-count:]]

    def source_title(self) -> str:
        pin = f"version v{self.version}" if self.version else "current release"
        return f"ONS time series {self.cdid} in dataset {self.dataset} ({pin}), JSON data endpoint"

    def point_item(self, point: dict, unit: str, period_word: str = "month") -> dict:
        snippet = f'"date":"{point["date"]}","value":"{point["value"]}","label":"{point["label"]}"'
        return cite(ONS_PUBLISHER, self.source_title(), self.url, self.raw, f"{period_word}s[date={point['date']}].value", point["value"], unit, snippet, OGL)

    def title_item(self) -> dict:
        """Evidence for the series identity; the title is the value unless JSON escaping keeps it from matching the raw bytes."""
        snippet = f'"datasetId":"{self.dataset}","datasetUri":"{self.description["datasetUri"]}","cdid":"{self.cdid}"'
        if collapse(self.title) in collapse(self.raw.decode("utf-8")):
            return cite(ONS_PUBLISHER, self.source_title(), self.url, self.raw, "description.title", self.title, "text", snippet, OGL)
        return cite(ONS_PUBLISHER, self.source_title(), self.url, self.raw, f"description.cdid (title: {self.title})", self.cdid, "text", snippet, OGL)

    def note(self) -> str:
        pin = f"pinned to published version v{self.version}, whose bytes do not change" if self.version else "read from the current endpoint, whose bytes change at each release"
        return (
            f"Series {self.cdid} ({self.title}) from ONS dataset {self.dataset}, {pin}; release date {self.description['releaseDate'][:10]}. "
            "The ONS JSON endpoint lists the series as months, quarters and years; only the periods in observed_period are kept. "
            "ONS publishes its data under the Open Government Licence v3.0, as the cpih01 dataset metadata states."
        )


def monthly_entry(series: OnsSeries, entry_id: str, title: str, domain: str, variable: str, unit: str, tags: list[str], use_cases: list[str], note: str, license_raw: bytes) -> dict:
    points = series.months(YEAR)
    if len(points) != 12:
        raise ValueError(f"{series.cdid}: expected 12 months for {YEAR}, found {len(points)}")
    first, last = points[0], points[-1]
    values = [v for _, v, _ in points]
    return {
        "id": entry_id,
        "title": title,
        "kind": "temporal",
        "domain": domain,
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "observed_period": {"start": f"{YEAR}-01-01", "end": f"{YEAR}-12-31"},
        "tags": tags + ["ons", "time-series"],
        "use_cases": use_cases,
        "spec": {"variable": variable, "unit": unit, "frequency": "monthly", "columns": [], "series": [{"period": p, "value": v} for p, v, _ in points]},
        "evidence": [series.point_item(first[2], unit), series.point_item(last[2], unit), series.title_item(), ons_license_item(license_raw)],
        "assumed": [],
        "tests": [
            {"statistic": f"value[{first[0]}]", "expected": first[1], "tolerance": 1e-9},
            {"statistic": f"value[{last[0]}]", "expected": last[1], "tolerance": 1e-9},
            {"statistic": "max", "expected": max(values), "tolerance": 1e-9},
            {"statistic": "min", "expected": min(values), "tolerance": 1e-9},
            {"statistic": "count", "expected": 12, "tolerance": 0},
        ],
        "notes": note + " " + series.note(),
    }


def quarterly_entry(series: OnsSeries, entry_id: str, title: str, domain: str, variable: str, unit: str, tags: list[str], use_cases: list[str], note: str, license_raw: bytes, columns: dict[str, OnsSeries] | None = None) -> dict:
    points = series.quarters(QUARTERS)
    by_period = {p: {"period": p, "value": v} for p, v, _ in points}
    for name, column in (columns or {}).items():
        column_points = {p: v for p, v, _ in column.quarters(len(column.data["quarters"]))}
        for p in by_period:
            by_period[p][name] = column_points[p]
    first, last = points[0], points[-1]
    values = [v for _, v, _ in points]
    start_year, start_q = first[0].split("-")
    end_year, end_q = last[0].split("-")
    items = [series.point_item(first[2], unit, "quarter"), series.point_item(last[2], unit, "quarter"), series.title_item()]
    for name, column in (columns or {}).items():
        column_last = column.quarters(1)[0]
        items += [column.point_item(column_last[2], unit, "quarter"), column.title_item()]
    items.append(ons_license_item(license_raw))
    tests = [
        {"statistic": f"value[{first[0]}]", "expected": first[1], "tolerance": 1e-9},
        {"statistic": f"value[{last[0]}]", "expected": last[1], "tolerance": 1e-9},
        {"statistic": "max", "expected": max(values), "tolerance": 1e-9},
        {"statistic": "count", "expected": len(points), "tolerance": 0},
    ]
    for name in columns or {}:
        tests.append({"statistic": f"column[{name},{last[0]}]", "expected": by_period[last[0]][name], "tolerance": 1e-9})
    column_note = ""
    if columns:
        column_note = " Columns: " + "; ".join(f"{name} is series {c.cdid} ({c.title})" for name, c in columns.items()) + "."
    return {
        "id": entry_id,
        "title": title,
        "kind": "temporal",
        "domain": domain,
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "observed_period": {"start": f"{start_year}-{QUARTER_START[start_q]}", "end": f"{end_year}-{QUARTER_END[end_q]}"},
        "tags": tags + ["ons", "time-series", "quarterly"],
        "use_cases": use_cases,
        "spec": {"variable": variable, "unit": unit, "frequency": "quarterly", "columns": list(columns or {}), "series": list(by_period.values())},
        "evidence": items,
        "assumed": [],
        "tests": tests,
        "notes": note + column_note + " " + series.note(),
    }


# ---------------------------------------------------------------- RPI average prices


def rpi_average_prices(raw: bytes, license_raw: bytes) -> dict:
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    titles, cdids = rows[0], rows[1]
    monthly = [r for r in rows if len(r[0]) == 8 and r[0][4] == " "]
    latest: dict[str, tuple[str, str, str]] = {}
    for column, (title, cdid) in enumerate(zip(titles, cdids)):
        if not re.search(r"ave(?:rage)? price", title, re.I):
            continue
        observed = [(r[0], r[column]) for r in monthly if r[column] != ""]
        if observed:
            latest[cdid] = (title, *observed[-1])
    period = max(p for _, p, _ in latest.values())
    table = {
        cdid: {"item": re.sub(r"^RPI\s*:?\s*Ave(?:rage)? [Pp]rice\s*[-:]\s*", "", title).strip().rstrip(","), "price_pence": float(price), "series_title": title}
        for cdid, (title, p, price) in sorted(latest.items())
        if p == period
    }
    year, month = period.split()
    month_start = f"{year}-{MONTHS[month]}-01"
    month_end = (date(int(year), int(MONTHS[month]), 28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    bread, milk = "CZOH", "CZNT"
    return {
        "id": f"gb-rpi-average-prices-{year}-{MONTHS[month]}",
        "title": f"UK average retail prices of selected items (RPI average prices), {date.fromisoformat(month_start):%B %Y}, pence",
        "kind": "rule",
        "domain": "retail",
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "valid_from": month_start,
        "valid_to": month_end.isoformat(),
        "tags": ["prices", "retail", "rpi", "average-prices", "reference-table", "ons"],
        "use_cases": ["retail-demo", "pricing", "sales-generation", "invoicing", "personal-finance"],
        "spec": {
            "implementation": "almanac.rules.lookup_table",
            "parameters": {"dataset": "MM23", "period": period, "unit": "pence", "rows": table},
            "vectors": [
                {"input": {"key": bread, "column": "price_pence"}, "expected": table[bread]["price_pence"], "origin": "source"},
                {"input": {"key": milk, "column": "price_pence"}, "expected": table[milk]["price_pence"], "origin": "source"},
                {"input": "CZXX", "expected": False, "origin": "derived"},
            ],
        },
        "evidence": [
            cite(ONS_PUBLISHER, "ONS consumer price inflation time series dataset MM23, version v138, CSV", MM23_URL, raw, f"row {period}, column {bread}", table[bread]["price_pence"], "decimal", table[bread]["series_title"], OGL),
            cite(ONS_PUBLISHER, "ONS consumer price inflation time series dataset MM23, version v138, CSV", MM23_URL, raw, f"row {period}, column {milk}", table[milk]["price_pence"], "decimal", table[milk]["series_title"], OGL),
            ons_license_item(license_raw),
        ],
        "assumed": [
            {"parameter": "rows[*].item", "value": "series title with the 'RPI: Ave price -' prefix removed", "method": "derived", "rationale": "The item label is the ONS series title minus its common prefix; series_title keeps the title verbatim."},
        ],
        "tests": [
            {"statistic": "vector_pass_fraction", "expected": 1, "tolerance": 0},
            {"statistic": "count", "expected": len(table), "tolerance": 0},
            {"statistic": "max[price_pence]", "expected": max(r["price_pence"] for r in table.values()), "tolerance": 1e-9},
        ],
        "notes": (
            f"Average retail prices in pence for the {len(table)} RPI average price series that carry a value for {period} in version v138 of the ONS "
            "MM23 consumer price inflation time series CSV (one row per CDID, keyed by the series identifier; the unit of each item, such as per pint, "
            f"per kg or per litre, is part of its title). {period} is the latest month with values for these series in that version; series that stopped "
            "earlier are left out. ONS publishes the dataset under the Open Government Licence v3.0, as the cpih01 metadata states."
        ),
    }


# ---------------------------------------------------------------- business demography


def workbook(raw: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def area_counts(sheet, codes: set[str]) -> dict[str, tuple[str, int]]:
    """Area code -> (name, 2024 count) for the rows whose code is wanted."""
    found = {}
    for row in sheet.iter_rows(values_only=True):
        if row[0] in codes:
            found[row[0]] = (str(row[1]).title(), int(row[2]))
    missing = codes - set(found)
    if missing:
        raise ValueError(f"{sheet.title}: areas not found {sorted(missing)}")
    return found


def demography_by_region(raw: bytes, license_raw: bytes) -> dict:
    wb = workbook(raw)
    codes = {"K02000001", "K03000001", "K04000001", "E92000001", "W92000004", "S92000003", "N92000002"} | {f"E120000{n:02d}" for n in range(1, 10)}
    births, deaths, active = (area_counts(wb[name], codes) for name in ("Table 1.1d", "Table 2.1d", "Table 3.1d"))
    rows = {}
    for code in sorted(codes):
        name, b = births[code]
        d, a = deaths[code][1], active[code][1]
        rows[code] = {"name": name, "births": b, "deaths": d, "active": a, "birth_rate_percent": round(100 * b / a, 2), "death_rate_percent": round(100 * d / a, 2)}
    uk = rows["K02000001"]
    return {
        "id": "gb-business-births-deaths-by-region-2024",
        "title": "UK business births, deaths and active enterprises by region and country, 2024, with birth and death rates",
        "kind": "rule",
        "domain": "company",
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "valid_from": "2024-01-01",
        "valid_to": "2024-12-31",
        "tags": ["business-demography", "births", "deaths", "by-region", "reference-table", "ons"],
        "use_cases": ["company-generation", "kyb-onboarding", "market-sizing", "credit-risk"],
        "spec": {
            "implementation": "almanac.rules.lookup_table",
            "parameters": {"year": 2024, "rows": rows},
            "vectors": [
                {"input": {"key": "K02000001", "column": "births"}, "expected": uk["births"], "origin": "source"},
                {"input": {"key": "E12000007", "column": "active"}, "expected": rows["E12000007"]["active"], "origin": "source"},
                {"input": {"key": "K02000001", "column": "birth_rate_percent"}, "expected": uk["birth_rate_percent"], "origin": "derived"},
                {"input": "E12000099", "expected": False, "origin": "derived"},
            ],
        },
        "evidence": [
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 1.1d", DEMOGRAPHY_URL, raw, "sheet Table 1.1d, row K02000001, column 2024", uk["births"], "count", "Table 1.1d - Count Of Births Of New Enterprises For 2024 by District, Counties And Unitary Authorities Within Region And Country", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 2.1d", DEMOGRAPHY_URL, raw, "sheet Table 2.1d, row K02000001, column 2024", uk["deaths"], "count", "Table 2.1d - Count Of Deaths Of New Enterprises For 2024 by District, Counties And Unitary Authorities Within Region And Country", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 3.1d", DEMOGRAPHY_URL, raw, "sheet Table 3.1d, row K02000001, column 2024", uk["active"], "count", "Table 3.1d - Count Of Active Enterprises For 2024 by District, Counties And Unitary Authorities Within Region And Country", OGL),
            ons_license_item(license_raw),
        ],
        "assumed": [
            {"parameter": "rows[*].birth_rate_percent", "value": "100 * births / active", "method": "derived", "rationale": "ONS publishes the counts; the birth rate is births as a percentage of active enterprises, rounded to two decimals."},
            {"parameter": "rows[*].death_rate_percent", "value": "100 * deaths / active", "method": "derived", "rationale": "Deaths as a percentage of active enterprises, rounded to two decimals."},
        ],
        "tests": [
            {"statistic": "vector_pass_fraction", "expected": 1, "tolerance": 0},
            {"statistic": "count", "expected": len(rows), "tolerance": 0},
            {"statistic": "max[births]", "expected": max(r["births"] for r in rows.values()), "tolerance": 0},
        ],
        "notes": (
            "Counts of enterprise births, deaths and active enterprises for 2024 from Tables 1.1d, 2.1d and 3.1d of the ONS Business Demography 2024 "
            "reference workbook (published 20 November 2025), kept for the United Kingdom, Great Britain, England and Wales, England, the nine English "
            "regions, Wales, Scotland and Northern Ireland, keyed by GSS area code; the workbook's local authority rows are dropped. Counts are control "
            "rounded to base 5 by ONS. Birth and death rates are derived here as counts over active enterprises. ONS publishes the workbook under the "
            "Open Government Licence v3.0."
        ),
    }


def division_counts(sheet, column: int) -> dict[str, tuple[str, int]]:
    """SIC 2007 division code -> (name, count) from a sheet whose labels read '10 : Manufacture of ...' at division level."""
    found = {}
    for row in sheet.iter_rows(values_only=True):
        label = str(row[0] or "")
        match = re.match(r"^(\d{2}) : (.+)$", label)
        if match:
            found[match.group(1)] = (match.group(2).strip(), int(row[column]))
    return found


def demography_by_industry(raw: bytes, license_raw: bytes) -> dict:
    wb = workbook(raw)
    year_column = 6
    births, deaths, active = (division_counts(wb[name], year_column) for name in ("Table 1.2", "Table 2.2", "Table 3.2"))
    if set(births) != set(deaths) or set(births) != set(active):
        raise ValueError("division sets differ between the births, deaths and active tables")
    rows = {}
    for code in sorted(births):
        name, b = births[code]
        d, a = deaths[code][1], active[code][1]
        rows[code] = {"name": name, "births": b, "deaths": d, "active": a}
        rows[code]["birth_rate_percent"] = round(100 * b / a, 2) if a else None
        rows[code]["death_rate_percent"] = round(100 * d / a, 2) if a else None
    sample = "10"
    return {
        "id": "gb-business-births-deaths-by-industry-2024",
        "title": "UK business births, deaths and active enterprises by SIC 2007 division, 2024, with birth and death rates",
        "kind": "rule",
        "domain": "company",
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "valid_from": "2024-01-01",
        "valid_to": "2024-12-31",
        "tags": ["business-demography", "births", "deaths", "sic-2007", "by-industry", "reference-table", "ons"],
        "use_cases": ["company-generation", "kyb-onboarding", "market-sizing", "credit-risk", "customer-segmentation"],
        "spec": {
            "implementation": "almanac.rules.lookup_table",
            "parameters": {"year": 2024, "classification": "SIC 2007 division (two digits)", "rows": rows},
            "vectors": [
                {"input": {"key": sample, "column": "births"}, "expected": rows[sample]["births"], "origin": "source"},
                {"input": {"key": sample, "column": "deaths"}, "expected": rows[sample]["deaths"], "origin": "source"},
                {"input": {"key": sample, "column": "active"}, "expected": rows[sample]["active"], "origin": "source"},
                {"input": "00", "expected": False, "origin": "derived", "note": "no division 00 exists"},
            ],
        },
        "evidence": [
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 1.2", DEMOGRAPHY_URL, raw, f"sheet Table 1.2, row '{sample} : {rows[sample]['name']}', column 2024", rows[sample]["births"], "count", "Table 1.2 - Count Of Births Of New Enterprises For 2019 To 2024 Standard Industrial Classification (SIC2007) Group", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 2.2", DEMOGRAPHY_URL, raw, f"sheet Table 2.2, row '{sample} : {rows[sample]['name']}', column 2024", rows[sample]["deaths"], "count", "Table 2.2 - Count Of Deaths Of Enterprises For 2019 To 2024 Standard Industrial Classification (SIC2007) Group", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 3.2", DEMOGRAPHY_URL, raw, f"sheet Table 3.2, row '{sample} : {rows[sample]['name']}', column 2024", rows[sample]["active"], "count", "Table 3.2", OGL),
            ons_license_item(license_raw),
        ],
        "assumed": [
            {"parameter": "rows[*].birth_rate_percent", "value": "100 * births / active", "method": "derived", "rationale": "ONS publishes the counts; the birth rate is births as a percentage of active enterprises, rounded to two decimals, null where no enterprise is active."},
            {"parameter": "rows[*].death_rate_percent", "value": "100 * deaths / active", "method": "derived", "rationale": "Deaths as a percentage of active enterprises, rounded to two decimals, null where no enterprise is active."},
        ],
        "tests": [
            {"statistic": "vector_pass_fraction", "expected": 1, "tolerance": 0},
            {"statistic": "count", "expected": len(rows), "tolerance": 0},
            {"statistic": "max[active]", "expected": max(r["active"] for r in rows.values()), "tolerance": 0},
        ],
        "notes": (
            "Counts of enterprise births, deaths and active enterprises for 2024 at SIC 2007 division level (two-digit codes) from Tables 1.2, 2.2 and 3.2 "
            "of the ONS Business Demography 2024 reference workbook (published 20 November 2025); the three-digit group rows beneath each division are "
            "dropped. Counts are control rounded to base 5 by ONS, so small divisions show zeros. Birth and death rates are derived here as counts over "
            "active enterprises. ONS publishes the workbook under the Open Government Licence v3.0."
        ),
    }


def survival_by_region(raw: bytes, license_raw: bytes) -> dict:
    wb = workbook(raw)
    sheet = wb["Table 4.1"]
    cohort = "2019"
    rows: dict[str, dict] = {}
    in_cohort = False
    for row in sheet.iter_rows(values_only=True):
        label = str(row[0] or "").strip()
        if re.fullmatch(r"\d{4}", label):
            in_cohort = label == cohort
            continue
        if in_cohort and label and row[1] is not None:
            record = {"births": int(row[1])}
            for years in range(1, 6):
                survivors, percent = row[2 * years], row[2 * years + 1]
                record[f"survival_{years}y"] = int(survivors)
                record[f"survival_{years}y_percent"] = round(float(percent), 1)
            rows[label] = record
    if "Total" not in rows or len(rows) < 13:
        raise ValueError(f"Table 4.1 cohort {cohort}: found {sorted(rows)}")
    total = rows["Total"]
    return {
        "id": "gb-business-survival-rates-by-region-2019-births",
        "title": "UK survival of enterprises born in 2019 by region and country: one- to five-year survival counts and rates",
        "kind": "rule",
        "domain": "company",
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "valid_from": "2019-01-01",
        "valid_to": "2024-12-31",
        "tags": ["business-demography", "survival", "by-region", "reference-table", "ons"],
        "use_cases": ["company-generation", "kyb-onboarding", "credit-risk", "market-sizing"],
        "spec": {
            "implementation": "almanac.rules.lookup_table",
            "parameters": {"birth_year": int(cohort), "rows": rows},
            "vectors": [
                {"input": {"key": "Total", "column": "births"}, "expected": total["births"], "origin": "source"},
                {"input": {"key": "Total", "column": "survival_5y"}, "expected": total["survival_5y"], "origin": "source"},
                {"input": {"key": "London", "column": "survival_1y"}, "expected": rows["London"]["survival_1y"], "origin": "source"},
                {"input": "Mars", "expected": False, "origin": "derived"},
            ],
        },
        "evidence": [
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 4.1", DEMOGRAPHY_URL, raw, f"sheet Table 4.1, cohort {cohort}, row Total, column Births", total["births"], "count", "Table 4.1", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 4.1", DEMOGRAPHY_URL, raw, f"sheet Table 4.1, cohort {cohort}, row Total, column 5-year survival", total["survival_5y"], "count", "5-year survival", OGL),
            cite(ONS_PUBLISHER, "ONS business demography reference table 2024 (XLSX), Table 4.1", DEMOGRAPHY_URL, raw, f"sheet Table 4.1, cohort {cohort}, row London, column 1-year survival", rows["London"]["survival_1y"], "count", "Business Demography 2024", OGL),
            ons_license_item(license_raw),
        ],
        "assumed": [
            {"parameter": "rows[*].survival_*y_percent", "value": "published percentage rounded to one decimal", "method": "derived", "rationale": "The workbook stores survival percentages at full floating-point precision; they are rounded to one decimal place here, the counts are copied unchanged."},
        ],
        "tests": [
            {"statistic": "vector_pass_fraction", "expected": 1, "tolerance": 0},
            {"statistic": "count", "expected": len(rows), "tolerance": 0},
            {"statistic": "max[survival_5y_percent]", "expected": max(r["survival_5y_percent"] for r in rows.values()), "tolerance": 1e-9},
            {"statistic": "min[survival_1y_percent]", "expected": min(r["survival_1y_percent"] for r in rows.values()), "tolerance": 1e-9},
        ],
        "notes": (
            f"Survival of enterprises born in {cohort} from Table 4.1 of the ONS Business Demography 2024 reference workbook (published 20 November 2025): "
            "births and the number and percentage still active one to five years later, for the nine English regions, Wales, Scotland, Northern Ireland "
            "and the UK total, keyed by the region name as the workbook prints it. The 2019 cohort is the latest with five-year survival. Counts are "
            "control rounded to base 5 by ONS. ONS publishes the workbook under the Open Government Licence v3.0."
        ),
    }


# ---------------------------------------------------------------- DESNZ road fuel prices


def fuel_prices(raw: bytes, api_raw: bytes) -> dict:
    text = raw.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader)
    columns = {"ulsd_pump_price_pence_per_litre": 2, "ulsp_duty_pence_per_litre": 3, "ulsd_duty_pence_per_litre": 4, "vat_percent": 5}
    points, lines = [], {}
    for row in reader:
        if not row or not row[0]:
            continue
        day = datetime.strptime(row[0], "%d/%m/%Y").date()
        if day.year != YEAR:
            continue
        point = {"period": day.isoformat(), "value": float(row[1])}
        point.update({name: float(row[i]) for name, i in columns.items()})
        points.append(point)
        lines[day.isoformat()] = ",".join(row)
    first, last = points[0], points[-1]
    peak = max(points, key=lambda p: p["value"])
    api = json.loads(api_raw)
    attachment = next(a for a in api["details"]["attachments"] if a["url"] == FUEL_CSV_URL)
    return {
        "id": f"gb-road-fuel-pump-prices-weekly-{YEAR}",
        "title": f"UK weekly road fuel pump prices, unleaded petrol and diesel, with duty and VAT rates, {YEAR}",
        "kind": "temporal",
        "domain": "energy",
        "region": "GB",
        "status": "verified",
        "tier": "A",
        "observed_period": {"start": first["period"], "end": last["period"]},
        "tags": ["fuel-prices", "petrol", "diesel", "pump-prices", "fuel-duty", "weekly", "desnz"],
        "use_cases": ["fleet-demo", "automotive-demo", "logistics", "pricing", "retail-demo"],
        "spec": {"variable": "ulsp_pump_price_pence_per_litre", "unit": "decimal", "frequency": "weekly", "columns": list(columns), "series": points},
        "evidence": [
            cite(DESNZ_PUBLISHER, "Weekly road fuel prices (CSV) 2018 to 2026, GOV.UK statistics attachment", FUEL_CSV_URL, raw, f"row Date={datetime.fromisoformat(first['period']).strftime('%d/%m/%Y')}, ULSP pump price column", row_value(lines[first["period"]], 1), "decimal", lines[first["period"]], OGL),
            cite(DESNZ_PUBLISHER, "Weekly road fuel prices (CSV) 2018 to 2026, GOV.UK statistics attachment", FUEL_CSV_URL, raw, f"row Date={datetime.fromisoformat(last['period']).strftime('%d/%m/%Y')}, ULSD pump price column", row_value(lines[last["period"]], 2), "decimal", lines[last["period"]], OGL),
            cite(DESNZ_PUBLISHER, "Weekly road fuel prices (CSV) 2018 to 2026, GOV.UK statistics attachment", FUEL_CSV_URL, raw, "header row, column 2", header[1], "text", header[2], OGL),
            cite(DESNZ_PUBLISHER, "GOV.UK content API item for /government/statistics/weekly-road-fuel-prices (JSON)", FUEL_API_URL, api_raw, "details.attachments[].url for the 2018 to 2026 CSV", attachment["url"], "text", f'"title":"{attachment["title"]}"', OGL),
        ],
        "assumed": [],
        "tests": [
            {"statistic": f"value[{first['period']}]", "expected": first["value"], "tolerance": 1e-9},
            {"statistic": f"value[{last['period']}]", "expected": last["value"], "tolerance": 1e-9},
            {"statistic": f"column[ulsd_pump_price_pence_per_litre,{last['period']}]", "expected": last["ulsd_pump_price_pence_per_litre"], "tolerance": 1e-9},
            {"statistic": "max", "expected": peak["value"], "tolerance": 1e-9},
            {"statistic": "count", "expected": len(points), "tolerance": 0},
        ],
        "notes": (
            f"Weekly UK average pump prices in pence per litre for ultra low sulphur unleaded petrol (the series value) and ultra low sulphur diesel, with "
            f"the duty rate in pence per litre and the VAT rate applying to each, for every week of {YEAR} in the DESNZ weekly road fuel prices CSV "
            f"(2018 to 2026) attached to the GOV.UK statistics page, public_updated_at {api['public_updated_at'][:10]}. Periods are the survey dates the "
            "file prints (Mondays). Rows for other years are dropped. GOV.UK content and DESNZ statistics are published under the Open Government Licence v3.0."
        ),
    }


def row_value(line: str, index: int) -> str:
    return line.split(",")[index]


# ---------------------------------------------------------------- Ofgem price cap


# (first day of the cap period, Ofgem document path, Direct Debit dual fuel level, typical consumption basis)
CAP_PERIODS = [
    ("2025-01-01", "2024-11/Summary_of_Changes_to_Energy_Price_Cap_1_January_to_31_March_2025.pdf", "Direct Debit £1,717 £1,738 £21 1.2%", 1738, "2023"),
    ("2025-04-01", "2025-02/Summary-of-changes-to-energy-price-cap-1-April-to-30-June-2025_1.pdf", "Direct Debit £1,738 £1,849 £111 6.4%", 1849, "2023"),
    ("2025-07-01", "2025-05/Summary%20of%20changes%20to%20energy%20price%20cap%201%20July%20to%2030%20September%202025_0.pdf", "Direct Debit £1,849 £1,720 -£129 -7%", 1720, "2023"),
    ("2025-10-01", "2025-08/summary-of-changes-to-energy-price-cap-1-october-to-31-december-2025.pdf", "Direct Debit £1,720 £1,755 £35 2%", 1755, "2023"),
    ("2026-01-01", "2025-11/Summary-of-changes-to-energy-price-cap-1-January-to-31-March-2026.pdf", "Direct Debit £1,755 £1,758 £3 0.2%", 1758, "2023"),
    ("2026-04-01", "2026-02/Summary-of-changes-to-energy-price-cap-1-April-to-30-June-2026.pdf", "Direct Debit £1,758 £1,641 -£117 -7%", 1641, "2023"),
    ("2026-07-01", "2026-06/Summary-of-changes-to-energy-price-cap-1-July-to-30-September-2026-revised-TDCV.pdf", "Direct Debit £1,477 £1,663 13%", 1663, "2026"),
    ("2026-10-01", "2026-08/Summary-of-changes-to-energy-price-cap-1-October-to-31-December-2026.pdf", "Direct Debit £1,663 £1,723 4%", 1723, "2026"),
]


def price_cap(copyright_raw: bytes) -> dict:
    steps, items = [], []
    for start, path, snippet, level, basis in CAP_PERIODS:
        url = f"{OFGEM_FILES}/{path}"
        raw = fetch(url)
        steps.append({"from": start, "value": level, "typical_consumption_values": basis})
        items.append(cite(OFGEM_PUBLISHER, f"Ofgem, summary of changes to the energy price cap for the period starting {start} (PDF)", url, raw, "table of cap levels by payment method, Direct Debit row, new level", f"£{level:,}", "currency:GBP", snippet, OGL))
    items.append(cite(OFGEM_PUBLISHER, "Ofgem copyright statement (HTML)", OFGEM_COPYRIGHT_URL, copyright_raw, "Crown copyright paragraph", "Open Government Licence", "text", "You may use Crown copyright information (not including logos) free of charge in any format or medium", OGL, role="supporting"))
    return {
        "id": "gb-energy-price-cap-levels-2025-2026",
        "title": "Great Britain energy price cap, dual fuel Direct Debit level at typical consumption, cap periods from January 2025 to December 2026",
        "kind": "rule",
        "domain": "energy",
        "region": "GB",
        "status": "verified",
        "tier": "B",
        "valid_from": CAP_PERIODS[0][0],
        "valid_to": "2026-12-31",
        "tags": ["energy-price-cap", "ofgem", "electricity", "gas", "step-schedule", "utilities"],
        "use_cases": ["utilities-demo", "household-generation", "personal-finance", "pricing"],
        "spec": {
            "implementation": "almanac.rules.step_schedule",
            "parameters": {"currency": "GBP", "unit": "GBP per year, dual fuel, paying by Direct Debit, at Ofgem's typical domestic consumption values", "steps": steps},
            "vectors": [
                {"input": {"on": CAP_PERIODS[0][0]}, "expected": CAP_PERIODS[0][3], "origin": "source"},
                {"input": {"on": "2025-09-30"}, "expected": CAP_PERIODS[2][3], "origin": "derived", "note": "last day of the July to September 2025 period"},
                {"input": {"on": CAP_PERIODS[-1][0]}, "expected": CAP_PERIODS[-1][3], "origin": "source"},
            ],
        },
        "evidence": items,
        "assumed": [],
        "tests": [
            {"statistic": "vector_pass_fraction", "expected": 1, "tolerance": 0},
            {"statistic": "count", "expected": len(steps), "tolerance": 0},
            {"statistic": "max", "expected": max(s["value"] for s in steps), "tolerance": 0},
            {"statistic": "min", "expected": min(s["value"] for s in steps), "tolerance": 0},
        ],
        "notes": (
            "The annual bill for a dual fuel household paying by Direct Debit at Ofgem's typical domestic consumption values, as Ofgem's summary of "
            "changes document for each quarterly cap period states it, in pounds per year; the cap itself limits unit rates and standing charges, and "
            "this figure is Ofgem's headline illustration of it. Each step is the first day of a cap period (January, April, July and October). Ofgem "
            "revised its typical domestic consumption values from 1 July 2026, so the levels from that date are on the 2026 values and are not "
            "comparable with earlier steps on the 2023 values; each step records which basis applies. Ofgem's copyright page states that its Crown "
            "copyright material may be used under the Open Government Licence."
        ),
    }


# ---------------------------------------------------------------- main


def main() -> int:
    license_raw = fetch(ONS_LICENSE_URL)
    labour = "employmentandlabourmarket"
    monthly = [
        (OnsSeries("businessindustryandtrade/retailindustry", "J4MC", "DRSI", 131), "gb-internet-sales-share-of-retail-monthly-2025", "Great Britain internet sales as a percentage of total retail sales, monthly 2025", "retail", "internet_sales_share_percent", "percent", ["retail", "e-commerce", "internet-sales", "seasonality"], ["retail-demo", "digital-channels", "demand-forecasting", "sales-generation"], "Share of retail sales made online, Great Britain, monthly for 2025, from the ONS retail sales inquiry."),
        (OnsSeries(f"{labour}/peopleinwork/earningsandworkinghours", "KAB9", "EMP", 131), "gb-average-weekly-earnings-total-pay-monthly-2025", "Great Britain average weekly earnings, whole economy, total pay, seasonally adjusted, monthly 2025", "employment", "average_weekly_earnings_total_pay_gbp", "currency:GBP", ["earnings", "wages", "awe", "pay"], ["payroll-demo", "hr-analytics", "personal-finance", "workforce-planning"], "Average weekly earnings in pounds (total pay excluding arrears, whole economy, seasonally adjusted), Great Britain, monthly for 2025."),
        (OnsSeries(f"{labour}/peopleinwork/employmentandemployeetypes", "AP2Z", "UNEM", 128), "gb-vacancies-per-100-employee-jobs-monthly-2025", "UK job vacancies per 100 employee jobs, all industries, monthly 2025", "employment", "vacancies_per_100_employee_jobs", "ratio", ["vacancies", "labour-demand", "hiring"], ["hr-analytics", "workforce-planning", "capacity-planning"], "Vacancies per 100 employee jobs, UK, all industries, monthly for 2025 (ONS labels the series as superseded by the job openings rate, N4JE, but continues to publish it)."),
        (OnsSeries(f"{labour}/peopleinwork/employmentandemployeetypes", "AP2Y", "UNEM", 130), "gb-job-vacancies-monthly-2025", "UK job vacancies, all industries, thousands, monthly 2025", "employment", "vacancies_thousands", "decimal", ["vacancies", "labour-demand", "hiring"], ["hr-analytics", "workforce-planning", "market-sizing"], "Number of job vacancies in thousands, UK, all industries, monthly for 2025, from the ONS Vacancy Survey."),
        (OnsSeries(f"{labour}/peoplenotinwork/redundancies", "BEIR", "LMS", 128), "gb-redundancy-rate-monthly-2025", "UK redundancy rate per 1,000 employees, seasonally adjusted, monthly 2025", "employment", "redundancy_rate_per_1000_employees", "ratio", ["redundancies", "labour", "attrition"], ["hr-analytics", "workforce-planning", "credit-risk"], "Redundancies per 1,000 employees (ILO definition, seasonally adjusted), UK, as labelled by month for 2025 in the ONS Labour Force Survey series."),
        (OnsSeries(f"{labour}/peoplenotinwork/unemployment", "MGSX", "LMS", 128), "gb-unemployment-rate-monthly-2025", "UK unemployment rate, aged 16 and over, seasonally adjusted, monthly 2025", "employment", "unemployment_rate_percent", "percent", ["unemployment", "labour"], ["hr-analytics", "credit-risk", "market-sizing", "person-generation"], "Unemployment rate (aged 16 and over, seasonally adjusted), UK, as labelled by month for 2025 in the ONS Labour Force Survey series."),
        (OnsSeries(f"{labour}/peopleinwork/employmentandemployeetypes", "LF24", "LMS", 128), "gb-employment-rate-monthly-2025", "UK employment rate, aged 16 to 64, seasonally adjusted, monthly 2025", "employment", "employment_rate_percent", "percent", ["employment", "labour"], ["hr-analytics", "workforce-planning", "person-generation"], "Employment rate (aged 16 to 64, seasonally adjusted), UK, as labelled by month for 2025 in the ONS Labour Force Survey series."),
        (OnsSeries(f"{labour}/peoplenotinwork/economicinactivity", "LF2S", "LMS", 128), "gb-economic-inactivity-rate-monthly-2025", "UK economic inactivity rate, aged 16 to 64, seasonally adjusted, monthly 2025", "employment", "economic_inactivity_rate_percent", "percent", ["economic-inactivity", "labour"], ["hr-analytics", "person-generation", "household-generation"], "Economic inactivity rate (aged 16 to 64, seasonally adjusted), UK, as labelled by month for 2025 in the ONS Labour Force Survey series."),
        (OnsSeries(f"{labour}/peoplenotinwork/outofworkbenefits", "BCJD", "UNEM", 131), "gb-claimant-count-monthly-2025", "UK claimant count, people claiming unemployment-related benefits, seasonally adjusted, thousands, monthly 2025", "employment", "claimant_count_thousands", "decimal", ["claimant-count", "benefits", "labour"], ["hr-analytics", "credit-risk", "market-sizing"], "People claiming unemployment-related benefits (Jobseeker's Allowance and Universal Credit claimants required to search for work) in thousands, UK, seasonally adjusted, monthly for 2025."),
    ]
    for series, entry_id, title, domain, variable, unit, tags, use_cases, note in monthly:
        write(monthly_entry(series, entry_id, title, domain, variable, unit, tags, use_cases, note, license_raw))

    gdp = "economy/grossdomesticproductgdp"
    write(quarterly_entry(OnsSeries(gdp, "NRJS", "UKEA", 44), "gb-household-saving-ratio-quarterly", "UK households' saving ratio, seasonally adjusted, latest 24 quarters", "economy", "household_saving_ratio_percent", "percent", ["saving-ratio", "households", "consumer"], ["personal-finance", "financial-planning", "banking-demo", "household-generation"], "Households' and NPISH saving ratio (saving as a percentage of total resources, current prices, seasonally adjusted), UK, quarterly, the latest 24 quarters in the pinned version.", license_raw))
    write(quarterly_entry(OnsSeries(gdp, "IHYQ", "PN2", 45), "gb-gdp-quarter-on-quarter-growth-quarterly", "UK GDP quarter-on-quarter growth, chained volume measure, seasonally adjusted, latest 24 quarters", "economy", "gdp_quarter_on_quarter_growth_percent", "percent", ["gdp", "growth", "macro"], ["market-sizing", "financial-planning", "demand-forecasting"], "Gross domestic product, quarter-on-quarter growth in percent (chained volume measure, seasonally adjusted), UK, quarterly, the latest 24 quarters in the pinned version.", license_raw))

    sector = "economy/nationalaccounts/uksectoraccounts"
    net = OnsSeries(sector, "LRWW", "PROF")
    columns = {
        "gross_rate_of_return_percent": OnsSeries(sector, "LRWV", "PROF"),
        "manufacturing_net_rate_of_return_percent": OnsSeries(sector, "LRYC", "PROF"),
        "services_net_rate_of_return_percent": OnsSeries(sector, "LRYQ", "PROF"),
    }
    write(quarterly_entry(net, "gb-pnfc-net-rate-of-return-quarterly", "UK private non-financial corporations' net rate of return, with gross, manufacturing and services rates, latest 24 quarters", "company", "pnfc_net_rate_of_return_percent", "percent", ["profitability", "rate-of-return", "corporations"], ["company-generation", "financial-reporting", "credit-risk", "loan-pricing"], "Net rate of return (net operating surplus over net capital employed, percent) of UK private non-financial corporations, quarterly, the latest 24 quarters, from the ONS profitability of UK companies release.", license_raw, columns))

    write(rpi_average_prices(fetch(MM23_URL), license_raw))

    demography_raw = fetch(DEMOGRAPHY_URL)
    write(demography_by_region(demography_raw, license_raw))
    write(demography_by_industry(demography_raw, license_raw))
    write(survival_by_region(demography_raw, license_raw))

    write(fuel_prices(fetch(FUEL_CSV_URL), fetch(FUEL_API_URL)))
    write(price_cap(fetch(OFGEM_COPYRIGHT_URL)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
